import io
import os
import shutil
import sys
import tempfile
import unittest

from flask import Flask
from openpyxl import load_workbook

BACKEND_DIR = os.path.dirname(os.path.dirname(__file__))
if BACKEND_DIR not in sys.path:
    sys.path.insert(0, BACKEND_DIR)

from routes.export_routes import export_bp
from services.spend_quality_assessment.excel_export import (
    build_complete_analysis_workbook,
)
from shared import db


def sample_assessment():
    return {
        "totalRows": 12,
        "dateSource": {
            "displayName": "Invoice Date",
            "sourceColumn": "Invoice date",
            "fallback": False,
        },
        "warnings": [{
            "code": "DATE_FALLBACK_USED",
            "severity": "warning",
            "message": "A date fallback was used in an earlier run.",
        }],
        "executiveSummary": {
            "rows": [{"key": "timePeriod", "label": "Time period", "text": "Two years of data"}],
        },
        "spendBreakdown": {
            "ltmSpend": 900.0,
            "currentFySpend": 800.0,
            "priorFySpend": 700.0,
            "yoyAbs": 100.0,
            "yoyPct": 14.3,
        },
        "supplierBreakdown": {"totalSuppliers": 3, "suppliersTo80Pct": 1},
        "columnFillRate": {
            "feasible": True,
            "columns": [{
                "columnName": "=Potential formula",
                "sourceColumn": "Spend Amount",
                "order": 1,
                "fillRate": 98.5,
                "spendCoverage": 99.0,
            }],
        },
        "spendBifurcation": {
            "feasible": True,
            "positiveSpend": 1_000.0,
            "positivePctOfNet": 111.1,
            "negativeSpend": -100.0,
            "negativePctOfNet": -11.1,
            "netSpend": 900.0,
        },
        "datePivot": {
            "feasible": True,
            "years": [2025, 2026],
            "months": ["Jan", "Feb"],
            "cells": {
                "2025": {"1": 100.0, "2": 200.0},
                "2026": {"1": 300.0, "2": 400.0},
            },
        },
        "paretoAnalysis": {
            "feasible": True,
            "thresholds": [80, 95],
            "totalDatasetSpend": 900.0,
            "metrics": {
                "80": {
                    "totalSpend": 720.0,
                    "transactionCount": 7,
                    "uniqueTransactions": 6,
                    "supplierCount": 1,
                },
                "95": {
                    "totalSpend": 855.0,
                    "transactionCount": 11,
                    "uniqueTransactions": 9,
                    "supplierCount": 3,
                },
            },
        },
    }


def sample_views():
    return [
        {
            "viewId": "supplier_ranking",
            "title": "Supplier Spend Ranking",
            "chartType": "hbar",
            "excludedRows": 1,
            "tableData": [{"Supplier Name": "=Unsafe text", "Total Spend (USD)": 700.0}],
        },
        {
            "viewId": "spend_over_time",
            "title": "Spend Over Time",
            "chartType": "bar",
            "tableData": {
                "monthly": [{"Year": 2026, "Month": "Jan", "Total Spend (USD)": 300.0}],
                "yearly": [{"Year": "2026", "Total Spend (USD)": 300.0}],
            },
        },
        {
            "viewId": "category_drilldown",
            "title": "Category Drill-Down",
            "chartType": "tree_pivot",
            "treeData": [{
                "level": "l1",
                "name": "Indirect",
                "totalSpend": 900.0,
                "percentOfParent": 100.0,
                "percentOfTotal": 100.0,
                "children": [{
                    "level": "l2",
                    "name": "IT",
                    "totalSpend": 900.0,
                    "percentOfParent": 100.0,
                    "percentOfTotal": 100.0,
                    "children": [],
                }],
            }],
        },
    ]


class CompleteAnalysisWorkbookTests(unittest.TestCase):
    def test_workbook_contains_each_assessment_cut_and_generated_view(self):
        workbook_bytes = build_complete_analysis_workbook(sample_assessment(), sample_views())
        workbook = load_workbook(io.BytesIO(workbook_bytes), data_only=False)

        self.assertEqual(
            [
                "Spend Summary",
                "Fill Rate Summary",
                "Spend Bifurcation",
                "Date-wise Distribution",
                "Spend Cut 80%",
                "Spend Cut 95%",
                "View - Supplier Spend Ranking",
                "View - Spend Over Time",
                "View - Category Drill-Down",
            ],
            workbook.sheetnames,
        )
        self.assertEqual("'=Potential formula", workbook["Fill Rate Summary"]["A4"].value)
        self.assertEqual("'=Unsafe text", workbook["View - Supplier Spend Ranking"]["A7"].value)
        self.assertIsInstance(workbook["Spend Bifurcation"]["B4"].value, (int, float))

        drilldown_values = [
            cell.value for row in workbook["View - Category Drill-Down"].iter_rows() for cell in row
        ]
        self.assertIn("Indirect", drilldown_values)
        self.assertIn("IT", drilldown_values)


class CompleteAnalysisExportRouteTests(unittest.TestCase):
    def setUp(self):
        self.original_sessions_dir = db.SESSIONS_DIR
        self.session_id = "complete-analysis-export"
        self.tmpdir = tempfile.mkdtemp(prefix="complete-analysis-export-")
        db.close_session_db(self.session_id)
        db.SESSIONS_DIR = self.tmpdir
        with db.get_session_lock(self.session_id):
            conn = db.get_session_db(self.session_id)
            db.set_meta(conn, "executive_summary", sample_assessment())
            db.set_meta(conn, "view_results", sample_views())

    def tearDown(self):
        try:
            db.delete_session(self.session_id)
        finally:
            db.SESSIONS_DIR = self.original_sessions_dir
            shutil.rmtree(self.tmpdir, ignore_errors=True)

    def _app(self):
        app = Flask(__name__)
        app.register_blueprint(export_bp, url_prefix="/api")
        return app

    def test_export_route_returns_workbook_with_download_headers(self):
        with self._app().test_client() as client:
            response = client.post(
                "/api/export/xlsx/complete-analysis",
                json={"sessionId": self.session_id},
            )

        self.assertEqual(200, response.status_code)
        self.assertTrue(response.content_type.startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ))
        self.assertEqual(
            'attachment; filename="spend-quality-and-views.xlsx"',
            response.headers["Content-Disposition"],
        )
        self.assertEqual("no-store", response.headers["Cache-Control"])
        self.assertIn("View - Category Drill-Down", load_workbook(io.BytesIO(response.data)).sheetnames)

    def test_export_route_requires_generated_views(self):
        with db.get_session_lock(self.session_id):
            db.set_meta(db.get_session_db(self.session_id), "view_results", [])

        with self._app().test_client() as client:
            response = client.post(
                "/api/export/xlsx/complete-analysis",
                json={"sessionId": self.session_id},
            )

        self.assertEqual(409, response.status_code)
        self.assertEqual("VIEWS_NOT_READY", response.get_json()["code"])
