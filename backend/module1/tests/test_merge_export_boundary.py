"""Regressions for the public schema of merge downloads and handoffs."""

from __future__ import annotations

import csv
import importlib
import io
import sys
import uuid
import zipfile

import pytest
from openpyxl import load_workbook


@pytest.fixture(scope="module")
def app_client():
    """Spin up the Module 1 Flask app using the isolated test sessions."""
    sys.modules.pop("app", None)
    app_module = importlib.import_module("app")
    app_module.app.config.update(TESTING=True)
    return app_module.app.test_client(), app_module


def _csv_rows(content: bytes) -> list[list[str]]:
    return list(csv.reader(io.StringIO(content.decode("utf-8"))))


def _assert_public_csv(content: bytes) -> None:
    rows = _csv_rows(content)
    assert rows == [["Invoice No.", "Amount"], ["INV-1", "42.50"]]
    assert "__row_id" not in content.decode("utf-8")
    assert "__source_table" not in content.decode("utf-8")


def _seed_merge_output(session_id: str) -> None:
    from shared.db import get_session_db, set_meta

    conn = get_session_db(session_id)
    for table_name in ("final_merged", "merge_v1", "_merge_step_source"):
        conn.execute(
            f'''CREATE TABLE "{table_name}" (
                "Invoice No." VARCHAR,
                "Amount" VARCHAR,
                "__row_id" BIGINT,
                "__source_table" VARCHAR
            )'''
        )
        conn.execute(
            f'''INSERT INTO "{table_name}" VALUES
                ('INV-1', '42.50', 1, 'legacy-source')'''
        )
    set_meta(conn, "merge_history", [{"version": 1, "table_name": "merge_v1", "file_label": "merge_v1"}])
    set_meta(conn, "groupSchemaTableRows", [{"group_id": "source", "group_name": "Source"}])
    conn.commit()


def test_merge_downloads_and_analyzer_transfer_share_a_public_schema(app_client, monkeypatch):
    """No serializer may pair a public SQL projection with internal headers/keys."""
    client, app_module = app_client
    session_id = uuid.uuid4().hex
    _seed_merge_output(session_id)
    query = {"sessionId": session_id}

    response = client.get("/api/merge/download-csv", query_string=query)
    assert response.status_code == 200
    _assert_public_csv(response.data)

    response = client.get("/api/merge/download-step-csv", query_string={**query, "sourceGroupId": "source"})
    assert response.status_code == 200
    _assert_public_csv(response.data)

    response = client.get("/api/merge/download-xlsx", query_string=query)
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.data), read_only=True)
    assert list(workbook.active.values) == [("Invoice No.", "Amount"), ("INV-1", "42.50")]

    response = client.get("/api/merge/download-step-xlsx", query_string={**query, "sourceGroupId": "source"})
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.data), read_only=True)
    assert list(workbook.active.values) == [("Invoice No.", "Amount"), ("INV-1", "42.50")]

    response = client.get("/api/merge/download-all-csv", query_string=query)
    assert response.status_code == 200
    with zipfile.ZipFile(io.BytesIO(response.data)) as archive:
        _assert_public_csv(archive.read("merge_v1.csv"))

    response = client.get("/api/merge/download-all", query_string=query)
    assert response.status_code == 200
    with zipfile.ZipFile(io.BytesIO(response.data)) as archive:
        workbook = load_workbook(io.BytesIO(archive.read("merge_v1")), read_only=True)
        assert list(workbook.active.values) == [("Invoice No.", "Amount"), ("INV-1", "42.50")]

    captured: dict[str, object] = {}

    class AnalyzerResponse:
        status_code = 200
        text = ""

        @staticmethod
        def json():
            return {"sessionId": "analyzer-session"}

    def fake_post(url, *, files, timeout):
        captured["url"] = url
        captured["csv"] = files["file"][1].read()
        captured["timeout"] = timeout
        return AnalyzerResponse()

    routes = importlib.import_module("routes.merging_routes")
    monkeypatch.setattr(routes._requests, "post", fake_post)
    app_module.app.config["MODULE3_BACKEND_URL"] = "http://analyzer.test"

    response = client.post("/api/merge/transfer-to-analyzer", json={"sessionId": session_id})
    assert response.status_code == 200
    assert response.get_json() == {"ok": True, "analyzerSessionId": "analyzer-session"}
    assert captured["url"] == "http://analyzer.test/api/import"
    _assert_public_csv(captured["csv"])
